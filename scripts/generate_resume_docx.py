import os

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.shared import Inches, Pt, RGBColor


def create_resume_docx(output_path):
    doc = Document()

    # Set page margins to 0.5 inch for professional ATS layout
    for section in doc.sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.5)
        section.right_margin = Inches(0.5)
        section.page_width = Inches(8.5)
        section.page_height = Inches(11.0)

    # Base font style
    normal_style = doc.styles["Normal"]
    normal_font = normal_style.font
    normal_font.name = "Calibri"
    normal_font.size = Pt(9.5)
    normal_font.color.rgb = RGBColor(0, 0, 0)

    def add_bottom_border(paragraph):
        pPr = paragraph._element.get_or_add_pPr()
        pBdr = parse_xml(
            r'<w:pBdr xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
            r'<w:bottom w:val="single" w:sz="6" w:space="1" w:color="000000"/>'
            r"</w:pBdr>"
        )
        pPr.append(pBdr)

    # 1. HEADER
    p_name = doc.add_paragraph()
    p_name.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_name.paragraph_format.space_before = Pt(0)
    p_name.paragraph_format.space_after = Pt(2)
    run_name = p_name.add_run("MARK REIZEL VINCENT C. PALMA")
    run_name.font.size = Pt(17)
    run_name.font.bold = True

    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(0)
    p_title.paragraph_format.space_after = Pt(3)
    run_title = p_title.add_run("FULL-STACK SOFTWARE ENGINEER")
    run_title.font.size = Pt(10)
    run_title.font.bold = True

    p_contact = doc.add_paragraph()
    p_contact.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_contact.paragraph_format.space_before = Pt(0)
    p_contact.paragraph_format.space_after = Pt(2)
    r_c = p_contact.add_run(
        "Las Piñas City, Metro Manila, Philippines   •   binssente@gmail.com   •   +63 915 805 7972"
    )
    r_c.font.size = Pt(9)

    p_portfolio = doc.add_paragraph()
    p_portfolio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_portfolio.paragraph_format.space_before = Pt(0)
    p_portfolio.paragraph_format.space_after = Pt(6)
    r_port = p_portfolio.add_run("Portfolio: https://bisente-portfolio.vercel.app")
    r_port.font.size = Pt(9)
    add_bottom_border(p_portfolio)

    def add_section_header(title):
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(3)
        run = p.add_run(title)
        run.font.size = Pt(10)
        run.font.bold = True
        add_bottom_border(p)

    def add_item_header(role, dates, company, location):
        table = doc.add_table(rows=2, cols=2)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        table.autofit = False
        table.columns[0].width = Inches(5.3)
        table.columns[1].width = Inches(2.2)

        # Row 0: Role and Date
        r0c0 = table.cell(0, 0).paragraphs[0]
        r0c0.paragraph_format.space_before = Pt(2)
        r0c0.paragraph_format.space_after = Pt(0)
        r0 = r0c0.add_run(role)
        r0.font.bold = True
        r0.font.size = Pt(9.5)

        r0c1 = table.cell(0, 1).paragraphs[0]
        r0c1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r0c1.paragraph_format.space_before = Pt(2)
        r0c1.paragraph_format.space_after = Pt(0)
        r1 = r0c1.add_run(dates)
        r1.font.bold = True
        r1.font.size = Pt(9)

        # Row 1: Company and Location
        r1c0 = table.cell(1, 0).paragraphs[0]
        r1c0.paragraph_format.space_before = Pt(0)
        r1c0.paragraph_format.space_after = Pt(2)
        r2 = r1c0.add_run(company)
        r2.font.italic = True
        r2.font.size = Pt(9)

        r1c1 = table.cell(1, 1).paragraphs[0]
        r1c1.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        r1c1.paragraph_format.space_before = Pt(0)
        r1c1.paragraph_format.space_after = Pt(2)
        r3 = r1c1.add_run(location)
        r3.font.size = Pt(9)

        # Remove borders
        for row in table.rows:
            for cell in row.cells:
                tcPr = cell._tc.get_or_add_tcPr()
                tcBorders = parse_xml(
                    r'<w:tcBorders xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">'
                    r'<w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/>'
                    r"</w:tcBorders>"
                )
                tcPr.append(tcBorders)

    def add_bullet(lead_bold, text):
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(1.5)
        p.paragraph_format.line_spacing = 1.15
        if lead_bold:
            r_bold = p.add_run(lead_bold)
            r_bold.font.bold = True
            r_bold.font.size = Pt(9)
        r_text = p.add_run(text)
        r_text.font.size = Pt(9)

    # 2. PROFESSIONAL SUMMARY
    add_section_header("PROFESSIONAL SUMMARY")
    p_sum = doc.add_paragraph()
    p_sum.paragraph_format.space_before = Pt(2)
    p_sum.paragraph_format.space_after = Pt(4)
    p_sum.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r_sum1 = p_sum.add_run("Full-Stack Software Engineer and ")
    r_sum1.font.size = Pt(9.5)
    r_sum2 = p_sum.add_run("Magna Cum Laude")
    r_sum2.font.bold = True
    r_sum2.font.size = Pt(9.5)
    r_sum3 = p_sum.add_run(
        " Information Systems graduate with 2+ years of production experience architecting end-to-end "
        "web applications, internal enterprise ERPs, and automated transaction platforms. Combines hands-on "
        "development in React, Next.js, Laravel, PHP, and PostgreSQL with a specialized Cisco cybersecurity "
        "background in network defense and threat management. Proven track record collaborating in small "
        "engineering teams and shipping resilient civic-tech software."
    )
    r_sum3.font.size = Pt(9.5)

    # 3. TECHNICAL SKILLS
    add_section_header("TECHNICAL SKILLS")
    skills = [
        ("Languages: ", "JavaScript (ES6+), TypeScript, PHP, Python, SQL, HTML5, CSS3"),
        (
            "Frontend: ",
            "React 19, Next.js, Vue 3, Vite, Tailwind CSS, Framer Motion, Bootstrap, Responsive UI/UX",
        ),
        (
            "Backend & APIs: ",
            "Laravel 11, Django, FastAPI, Node.js, RESTful APIs, WebSockets (Socket.io), Sanctum, PHPMailer",
        ),
        (
            "Databases & BaaS: ",
            "PostgreSQL, MySQL, Supabase, Redis, AWS S3, Prisma ORM, Query Optimization",
        ),
        (
            "Cloud & Hosting: ",
            "Vercel, Render, Hostinger, Docker & Docker Compose, Linux, Nginx, CI/CD Actions, Git/GitHub",
        ),
        (
            "Cybersecurity & Systems: ",
            "SOC Operations, Incident Response, Network Defense, Threat Intelligence (CVSS), Endpoint Hardening, Cisco IOS",
        ),
        (
            "AI & Tools: ",
            "Claude, Google Gemini, CodeRabbit, Groq AI, Microsoft AutoGen",
        ),
    ]
    for cat, items in skills:
        add_bullet(cat, items)

    # 4. PROFESSIONAL EXPERIENCE
    add_section_header("PROFESSIONAL EXPERIENCE")

    add_item_header(
        "Software Development Associate",
        "June 2026 – Present",
        "AAA and Co., CPAs",
        "Metro Manila, Philippines",
    )
    add_bullet(
        "",
        "Architect and maintain the Gamma Oracle Project Management & Cost Control platform, streamlining internal audit workflows, financial reporting, and project cost tracking for accounting teams.",
    )
    add_bullet(
        "",
        "Implement atomic database transactions and strict access controls to prevent data discrepancy across multi-tiered corporate ledger allocations.",
    )
    add_bullet(
        "",
        "Collaborate with certified accountants and senior partners to translate complex statutory compliance requirements into automated operational dashboards.",
    )

    add_item_header(
        "Software Engineering Intern",
        "February 2026 – June 2026",
        "Gamma Oracle Dimension Inc.",
        "Metro Manila, Philippines",
    )
    add_bullet(
        "",
        "Co-developed the official educational web portal for Tax Leaders Circle PH (TLCPH.online) alongside a fellow developer intern, servicing accredited accounting seminars nationwide.",
    )
    add_bullet(
        "",
        "Engineered an automated certificate generator using FPDI/FPDF and a digital tax publication reader using PDF.js and Turn.js with realistic flip-book navigation.",
    )
    add_bullet(
        "",
        "Initiated the core development of BizMaker HRIS ERP, transitioning payroll workflows into a scalable cloud SaaS application with automated tax calculations and timesheet synchronization.",
    )

    add_item_header(
        "Freelance Full-Stack Developer",
        "January 2024 – Present",
        "Self-Employed / Independent Contractor",
        "Las Piñas City, Philippines",
    )
    add_bullet(
        "Bubble Hideout POS: ",
        "Engineered a standalone restaurant point-of-sale and live inventory management system in PHP/MySQL, eliminating dinner-rush order latency and automating daily cash-out reports.",
    )
    add_bullet(
        "Michael Anthony's Fashion Inc. (MAFI): ",
        "Automated manual paper invoices and Delivery Receipts (DR) into structured digital records, cutting invoicing errors and administrative overhead.",
    )
    add_bullet(
        "Galaxent Perfume: ",
        "Designed and deployed an attendance logging and payroll calculation portal for retail personnel, reducing manual bi-weekly time reconciliation.",
    )
    add_bullet(
        "STI OMMA Platform: ",
        "Built a production-ready community showcase portal for STI College multimedia artists featuring AWS S3 media uploads, Xendit payment gateways, and real-time chat.",
    )

    # 5. KEY PROJECTS & HACKATHONS
    doc.add_page_break()
    add_section_header("KEY PROJECTS & HACKATHONS")

    add_item_header(
        "eGovPH National Hackathon 2026 — Team Lead / Developer",
        "July 2026",
        "Department of Information and Communications Technology (DICT)",
        "SMX Aura, BGC",
    )
    add_bullet(
        "",
        'Led 5-member team "Println New Gen" in a 48-hour civic-tech sprint developing verification prototypes atop official government infrastructure (eGov Single Sign-On, eVerify, and eGovPay APIs).',
    )
    add_bullet(
        "",
        "Designed client-facing authentication flows and interfaced live with sandbox government identity verification endpoints.",
    )

    add_item_header(
        "MAFI 3D Custom Clothing Platform — Capstone Project Lead Developer",
        "2025",
        "Michael Anthony's Fashion Inc. / DFCAMCLP",
        "Las Piñas City",
    )
    add_bullet(
        "",
        "Led a year-long capstone converting a traditional garment manufacturer's manual sketching bottleneck into a browser-based 3D visual tailoring customizer.",
    )
    add_bullet(
        "",
        "Allowed real-time collar, fabric, and trim configurations that generate production-ready specifications directly for pattern cutters.",
    )

    add_item_header(
        "CND Upraze — Lead Full-Stack Architect and Developer",
        "2026",
        "Technopreneurship",
        "Academic Project",
    )
    add_bullet(
        "",
        "Architected an enterprise SaaS application uniting a React 19 frontend and a decoupled Laravel 11 protocol core, packaged in Docker for isolated local and staging environments.",
    )
    add_bullet(
        "",
        "Integrated Groq AI for automated inference and telemetry visualization dashboards for administrative order governance.",
    )

    # 6. EDUCATION & LEADERSHIP
    add_section_header("EDUCATION & LEADERSHIP")

    add_item_header(
        "Bachelor of Science in Information Systems (BSIS)",
        "2022 – 2026",
        "Dr. Filemon C. Aguilar Memorial College of Las Piñas (DFCAMCLP)",
        "Magna Cum Laude",
    )
    add_bullet(
        "Graphic Artist Society: ",
        "Elected Auditor (2025 – 2026); served as Level Representative (2023 – 2025), coordinating campus creative workshops and managing organizational inventory.",
    )

    add_item_header(
        "High School Diploma — TVL-ICT Track",
        "2016 – 2022",
        "De La Salle Santiago Zobel School – BRafeNHS",
        "Muntinlupa City",
    )
    add_bullet(
        "",
        "Senior High School Technical-Vocational-Livelihood (TVL-ICT Track, 2020–2022)   •   Junior High School (2016–2020).",
    )

    # 7. CERTIFICATIONS & ACCREDITATIONS
    add_section_header("CERTIFICATIONS & ACCREDITATIONS")
    certs = [
        (
            "Junior Cybersecurity Analyst Career Path",
            " — Cisco Networking Academy (2026)",
        ),
        ("Cyber Threat Management", " — Cisco Networking Academy (2026)"),
        ("Network Defense", " — Cisco Networking Academy (2026)"),
        (
            "Endpoint Security & Introduction to Cybersecurity",
            " — Cisco Networking Academy (2026)",
        ),
        (
            "Networking Devices & Configuration • Network Support",
            " — Cisco Networking Academy (2026)",
        ),
        (
            "AI Fundamentals: Foundations for Understanding AI",
            " — IBM SkillsBuild & Cisco (2026)",
        ),
        (
            "Data Science Essentials with Python • Data Analytics Essentials",
            " — Cisco (2026)",
        ),
        (
            "Python Essentials (1 & 2) • JavaScript Essentials (1 & 2)",
            " — Cisco & OpenEDG (2026)",
        ),
        (
            "Intellectual Property Rights for Freelancers",
            " — DICT Region 2 (2026)   •   Cyber-Shield — DICT Region 8 (2026)",
        ),
    ]
    for c_title, c_meta in certs:
        add_bullet(c_title, c_meta)

    doc.save(output_path)
    print(f"Successfully generated non-HTML DOCX: {output_path}")


if __name__ == "__main__":
    out = os.path.join(
        os.path.dirname(__file__),
        "..",
        "public",
        "Mark_Reizel_Vincent_Palma_Resume.docx",
    )
    out = os.path.abspath(out)
    create_resume_docx(out)
